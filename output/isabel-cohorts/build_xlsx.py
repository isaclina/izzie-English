import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.comments import Comment

d=json.load(open('data.json'))
F='Arial'
BLUE=Font(name=F,color='0000FF'); BLK=Font(name=F); GRN=Font(name=F,color='008000')
BOLD=Font(name=F,bold=True); TITLE=Font(name=F,bold=True,size=14); MUTED=Font(name=F,color='666666',size=9)
HDR=PatternFill('solid',fgColor='DDE7E4'); AVG=PatternFill('solid',fgColor='F2F2F2')
thin=Side(style='thin',color='BFBFBF'); BOX=Border(top=thin,bottom=thin,left=thin,right=thin)
C=Alignment(horizontal='center',vertical='center',wrap_text=True)
GO=['중1','중2','중3']

wb=Workbook(); ws0=wb.active; ws0.title='안내'
def cell(ws,r,c,v,font=BLK,fill=None,fmt=None,align=C):
    x=ws.cell(r,c,v); x.font=font; x.border=BOX; x.alignment=align
    if fill: x.fill=fill
    if fmt: x.number_format=fmt
    return x

refs={}  # cohort -> dict of avg cell refs per exam col key
sheetnames={'A':'김건희반_2020','B':'구성관반_2023','C':'이서연반_2025'}
for co in 'ABC':
    info=d['cohorts'][co]; ws=wb.create_sheet(sheetnames[co])
    rows=[r for r in d['mock'] if r['cohort']==co]
    studs=list(dict.fromkeys(r['student'] for r in rows))
    exams=sorted({(r['grade'],r['seq'],r['level']) for r in rows},key=lambda e:(GO.index(e[0]),e[1]))
    get={(r['student'],r['grade'],r['seq']):r for r in rows}
    ws['A1']=f"{info['name']} · 모의고사 성적 ({info['years']})"; ws['A1'].font=TITLE
    ws['A2']=f"반: {info['classes']} · 파란 글씨 = 성적표 사진에서 옮긴 입력값, 검은 글씨 = 계산식. 빈칸 = 기록 없음/미응시."; ws['A2'].font=MUTED
    n=len(exams); c0=2; cl=c0+n-1  # exam columns B..
    def header(r,label):
        cell(ws,r,1,label,BOLD,HDR)
        for j,(g,q,lv) in enumerate(exams):
            cell(ws,r,c0+j,f"{g} {q}차\n{lv.replace(' 모의고사','').replace('2023학년도 ','')}",BOLD,HDR)
    r=4; blocks={}
    for key,label,extra in [('total','총점',True),('tier','등급',False),('listening','듣기 (맞은 개수 / 17)',False),('reading','독해 (맞은 개수 / 28)',False)]:
        cell(ws,r,1,label,BOLD); ws.cell(r,1).border=Border(); r+=1
        header(r,'학생')
        if extra:
            for k,h in enumerate(['첫 시험','마지막 시험','변화','최고']): cell(ws,r,cl+1+k,h,BOLD,HDR)
        hr=r; r+=1; first=r
        for s in studs:
            cell(ws,r,1,s,BOLD)
            for j,(g,q,lv) in enumerate(exams):
                v=get.get((s,g,q),{}).get(key)
                x=cell(ws,r,c0+j,v,BLUE)
                note=get.get((s,g,q),{}).get('note')
                if v is None and note and key=='total': x.comment=Comment(note,'Isabel')
            if extra:
                rng=f"{L(c0)}{r}:{L(cl)}{r}"
                cell(ws,r,cl+1,f"={L(c0)}{r}")
                cell(ws,r,cl+2,f"=LOOKUP(9.99E+307,{rng})")
                cell(ws,r,cl+3,f"={L(cl+2)}{r}-{L(cl+1)}{r}",fmt='+0;-0;0')
                cell(ws,r,cl+4,f"=MAX({rng})")
            r+=1
        last=r-1
        if key!='tier':
            cell(ws,r,1,'평균',BOLD,AVG)
            for j in range(n):
                col=L(c0+j)
                cell(ws,r,c0+j,f'=IF(COUNT({col}{first}:{col}{last})=0,"",AVERAGE({col}{first}:{col}{last}))',BLK,AVG,'0.0')
            blocks[key]=(hr,first,last,r); r+=1
        else: blocks[key]=(hr,first,last,None)
        r+=1
    ws.column_dimensions['A'].width=22
    for j in range(n+4): ws.column_dimensions[L(c0+j)].width=10
    for rr in range(1,r): pass
    for (hr,_,_,_) in blocks.values(): ws.row_dimensions[hr].height=32
    # 내신
    nrows=[x for x in d['naesin'] if x['cohort']==co]
    terms=list(dict.fromkeys((x['grade'],x['year'],x['term']) for x in nrows))
    cell(ws,r,1,'학교 내신 (영어)',BOLD); ws.cell(r,1).border=Border(); r+=1
    cell(ws,r,1,'학생',BOLD,HDR)
    for j,(g,y,t) in enumerate(terms): cell(ws,r,2+j,f"{g} ({y})\n{t}",BOLD,HDR)
    cell(ws,r,2+len(terms),'평균',BOLD,HDR); ws.row_dimensions[r].height=32; r+=1
    ng={(x['student'],x['grade'],x['term']):x for x in nrows}; nf=r
    for s in studs:
        cell(ws,r,1,s,BOLD)
        for j,(g,y,t) in enumerate(terms):
            x=ng.get((s,g,t)); v=x['score'] if x else None
            c=cell(ws,r,2+j,v,BLUE)
            if x and x.get('note'): c.comment=Comment(x['note'],'Isabel')
        rng=f"B{r}:{L(1+len(terms))}{r}"
        cell(ws,r,2+len(terms),f'=IF(COUNT({rng})=0,"",AVERAGE({rng}))',fmt='0.0'); r+=1
    r+=1
    # chart
    hr,first,last,avg=blocks['total']
    ch=LineChart(); ch.title='총점 추이'; ch.y_axis.title='총점'; ch.y_axis.scaling.min=40; ch.y_axis.scaling.max=100
    ch.height=9; ch.width=max(18,1.6*n)
    data=Reference(ws,min_col=1,max_col=cl,min_row=first,max_row=last)
    ch.add_data(data,from_rows=True,titles_from_data=True)
    ch.set_categories(Reference(ws,min_col=c0,max_col=cl,min_row=hr,max_row=hr))
    for sr in ch.series: sr.marker.symbol='circle'; sr.smooth=False
    ws.add_chart(ch,f"A{r+1}")
    ws.freeze_panes='B4'
    refs[co]=dict(ws=ws.title,exams=exams,c0=c0,blocks=blocks)

# 반 비교
wc=wb.create_sheet('반 비교',1)
wc['A1']='반별 같은 단계 비교 (평균)'; wc['A1'].font=TITLE
wc['A2']='초록 글씨 = 각 반 시트의 평균 행을 가져온 값. 중2 3차는 세 반 모두 처음 본 고2 모의고사입니다.'; wc['A2'].font=MUTED
stages=[('중1 첫 시험 (고1)',lambda ex:ex.index(next(e for e in ex if e[0]=='중1'))),
        ('중1 마지막 시험',lambda ex:max(i for i,e in enumerate(ex) if e[0]=='중1')),
        ('중2 첫 시험',lambda ex:ex.index(next(e for e in ex if e[0]=='중2'))),
        ('첫 고2 모의고사 (중2 3차)',lambda ex:ex.index(next(e for e in ex if e[0]=='중2' and e[1]==3))),
        ('가장 최근 시험',lambda ex:len(ex)-1)]
cell(wc,4,1,'총점 평균',BOLD,HDR)
for k,co in enumerate('ABC'): cell(wc,4,2+k,d['cohorts'][co]['short'],BOLD,HDR)
cell(wc,12,1,'독해 평균 (맞은 개수)',BOLD,HDR)
for k,co in enumerate('ABC'): cell(wc,12,2+k,d['cohorts'][co]['short'],BOLD,HDR)
for i,(lab,fn) in enumerate(stages):
    cell(wc,5+i,1,lab,BOLD); cell(wc,13+i,1,lab,BOLD)
    for k,co in enumerate('ABC'):
        R=refs[co]; j=fn(R['exams']); col=L(R['c0']+j); e=R['exams'][j]
        cell(wc,5+i,2+k,f"='{R['ws']}'!{col}{R['blocks']['total'][3]}",GRN,fmt='0.0').comment=Comment(f"{e[0]} {e[1]}차 {e[2]}",'Isabel')
        cell(wc,13+i,2+k,f"='{R['ws']}'!{col}{R['blocks']['reading'][3]}",GRN,fmt='0.0')
cell(wc,10,1,'중1 첫 시험 → 첫 고2 모의고사 변화',BOLD,AVG)
for k in range(3): cell(wc,10,2+k,f"={L(2+k)}8-{L(2+k)}5",BLK,AVG,'+0.0;-0.0;0.0')
wc.column_dimensions['A'].width=34
for k in range(3): wc.column_dimensions[L(2+k)].width=18
bc=BarChart(); bc.type='col'; bc.title='단계별 반 평균 총점'; bc.y_axis.scaling.min=50; bc.y_axis.scaling.max=100; bc.height=9; bc.width=20
bc.add_data(Reference(wc,min_col=2,max_col=4,min_row=4,max_row=9),titles_from_data=True)
bc.set_categories(Reference(wc,min_col=1,min_row=5,max_row=9))
wc.add_chart(bc,'F4')

# 안내
lines=[('이사벨 반 모의고사·내신 성적 분석',TITLE),
 ('Metta 영어 전문학원 · 성적표 사진(2020–2026)에서 옮긴 자료',MUTED),('',BLK),
 ('시트 구성',BOLD),
 ('· 반 비교: 세 반을 같은 단계(중1 첫 시험, 첫 고2 모의고사 등)에서 비교',BLK),
 ('· 김건희반_2020: 2020 중1 상위 4명, 중1~중3 (2020–2022), 중3 내신',BLK),
 ('· 구성관반_2023: 2023 중1 반 4명, 중1~중3 2차 (2023–2025), 중1·중2 내신',BLK),
 ('· 이서연반_2025: 2025 중1 반 중 2026년까지 다닌 5명, 중1 1차 + 중2 1~3차, 중2 내신',BLK),('',BLK),
 ('읽는 법',BOLD),
 ('· 파란 글씨: 성적표에서 옮긴 값 (수정 가능) · 검은 글씨: 계산식 · 초록 글씨: 다른 시트에서 가져온 값',BLK),
 ('· 듣기 17문항, 독해 28문항 기준 맞은 개수. 총점 100점, 절대평가 등급(90점 이상 1등급, 10점 단위)',BLK),
 ('· 빈칸은 성적표에 기록이 없거나 미응시. 메모(빨간 삼각형)에 이유를 적어 두었습니다.',BLK),('',BLK),
 ('참고',BOLD),
 ('· 김건희반 중2 1·2차는 성적표에 시험 수준(고1/고2) 표기가 없습니다.',BLK),
 ('· 구성관반 내신은 학기 표기가 없어 "중간·기말"로만 적었습니다. 2025년 3·4차(고3)는 받은 사진에 없습니다.',BLK),
 ('· 이서연반 2025년 중1은 1차만 사진에 있습니다. 2026년 이시원·한재승(중2 합류), 2025년 김은혜·구승준·김가민(중2 미등록)은 제외했습니다.',BLK),
 ('· 구성관반 김의현·최리원(2023 1차만 응시), 백시우·신상훈(2025 합류)은 제외했습니다.',BLK)]
for i,(t,f) in enumerate(lines,1): ws0.cell(i,1,t).font=f
ws0.column_dimensions['A'].width=110
wb.save('이사벨반_성적분석.xlsx'); print('saved')
