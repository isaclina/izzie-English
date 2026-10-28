import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment

import sys
d = json.load(open('data.json'))
CO = sys.argv[1]
NAME = {'A': '김건희 반', 'B': '구성관 반', 'C': '이서연 반'}[CO]
rows = [r for r in d['mock'] if r['cohort'] == CO]
STUDS = list(dict.fromkeys(r['student'] for r in rows))
get = {(r['student'], r['grade'], r['seq']): r for r in rows}
GO = ['중1', '중2', '중3']
def lvl(r):
    lv = r['level']
    if '고2' in lv: return '고2 모의고사'
    if '고3' in lv or '수능' in lv: return '고3 모의고사·수능'
    return '고1 모의고사'
allex = sorted({(r['grade'], r['seq'], r['year'], lvl(r)) for r in rows}, key=lambda e: (GO.index(e[0]), e[1]))
GROUPS = []
for g in ['고1 모의고사', '고2 모의고사', '고3 모의고사·수능']:
    ex = [e for e in allex if e[3] == g]
    if not ex: continue
    parts = []
    for gr in GO:
        qs = [e[1] for e in ex if e[0] == gr]
        if qs: parts.append(f"{gr}({next(e[2] for e in ex if e[0] == gr)}) {'·'.join(map(str, qs))}차")
    GROUPS.append((g, ', '.join(parts), [(e[0], e[1]) for e in ex]))
MISSING = [g for g in ['고1 모의고사', '고2 모의고사', '고3 모의고사·수능'] if g not in [x[0] for x in GROUPS]]
METRICS = [('total', '총점', '점 / 100', 40, 100, 10), ('listening', '듣기', '맞은 개수 / 17', 10, 17, 1), ('reading', '독해', '맞은 개수 / 28', 10, 28, 2)]

F = 'Arial'
BLUE = Font(name=F, color='0000FF'); BLK = Font(name=F); GRN = Font(name=F, color='008000')
BOLD = Font(name=F, bold=True); TITLE = Font(name=F, bold=True, size=14); MUTED = Font(name=F, color='666666', size=9)
HDR = PatternFill('solid', fgColor='DDE7E4'); AVG = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='BFBFBF'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)

def cell(ws, r, c, v, font=BLK, fill=None, fmt=None):
    x = ws.cell(r, c, v); x.font = font; x.border = BOX; x.alignment = CEN
    if fill: x.fill = fill
    if fmt: x.number_format = fmt
    return x

def label(g, q):
    r = next(x for x in rows if x['grade'] == g and x['seq'] == q)
    lv = r['level'].replace('2023학년도 ', '')
    if lv == '모의고사': lv = '수준 표기 없음'
    return f"{g} {q}차\n{lv}"

wb = Workbook(); info = wb.active; info.title = '안내'
summary = []  # (group, metric, sheet, avg_row, first_col, last_col)
for gname, gdesc, exams in GROUPS:
    ws = wb.create_sheet(gname.replace('·', '_'))
    ws['A1'] = f'{NAME} · {gname} 성적 변화'; ws['A1'].font = TITLE
    ws['A2'] = f'{gdesc} · 파란 글씨 = 성적표 입력값, 검은 글씨 = 계산식, 빈칸 = 기록 없음/미응시'; ws['A2'].font = MUTED
    n = len(exams); c0 = 2; cl = c0 + n - 1
    r = 4; chart_row = 4
    for key, name, unit, lo, hi, step in METRICS:
        ws.cell(r, 1, f'{name} ({unit})').font = BOLD; r += 1
        hr = r
        cell(ws, r, 1, '학생', BOLD, HDR)
        for j, (g, q) in enumerate(exams): cell(ws, r, c0 + j, label(g, q), BOLD, HDR)
        for k, h in enumerate(['첫 시험', '마지막 시험', '변화']): cell(ws, r, cl + 1 + k, h, BOLD, HDR)
        ws.row_dimensions[r].height = 34; r += 1; first = r
        for s in STUDS:
            cell(ws, r, 1, s, BOLD)
            fc = None
            for j, (g, q) in enumerate(exams):
                x = get.get((s, g, q), {}); v = x.get(key)
                c = cell(ws, r, c0 + j, v, BLUE)
                if v is None and x.get('note'): c.comment = Comment(x['note'], 'Isabel')
                if v is not None and fc is None: fc = c0 + j
            rng = f'{L(c0)}{r}:{L(cl)}{r}'
            if fc is None:
                for k in range(3): cell(ws, r, cl + 1 + k, None)
            else:
                cell(ws, r, cl + 1, f'={L(fc)}{r}')
                cell(ws, r, cl + 2, f'=LOOKUP(9.99E+307,{rng})')
                cell(ws, r, cl + 3, f'={L(cl+2)}{r}-{L(cl+1)}{r}', fmt='+0;-0;0')
            r += 1
        last = r - 1
        cell(ws, r, 1, '반 평균', BOLD, AVG)
        for j in range(n):
            col = L(c0 + j)
            cell(ws, r, c0 + j, f'=IF(COUNT({col}{first}:{col}{last})=0,"",AVERAGE({col}{first}:{col}{last}))', BLK, AVG, '0.0')
        fa = next(j for j in range(n) if any(get.get((s, *exams[j]), {}).get(key) is not None for s in STUDS))
        cell(ws, r, cl + 1, f'={L(c0+fa)}{r}', BLK, AVG, '0.0')
        cell(ws, r, cl + 2, f'=LOOKUP(9.99E+307,{L(c0)}{r}:{L(cl)}{r})', BLK, AVG, '0.0')
        cell(ws, r, cl + 3, f'={L(cl+2)}{r}-{L(cl+1)}{r}', BLK, AVG, '+0.0;-0.0;0.0')
        summary.append((gname, name, ws.title, r, cl + 1, cl + 2, cl + 3))
        # chart
        ch = LineChart(); ch.title = f'{gname} · {name}'; ch.y_axis.title = unit; ch.x_axis.title = None
        ch.y_axis.scaling.min = lo; ch.y_axis.scaling.max = hi; ch.y_axis.majorUnit = step
        ch.y_axis.delete = False; ch.x_axis.delete = False
        ch.height = 7.5; ch.width = 16; ch.display_blanks = 'gap'
        ch.add_data(Reference(ws, min_col=1, max_col=cl, min_row=first, max_row=last), from_rows=True, titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=c0, max_col=cl, min_row=hr, max_row=hr))
        for sr in ch.series: sr.marker.symbol = 'circle'; sr.marker.size = 7; sr.smooth = False
        ws.add_chart(ch, f'{L(cl + 5)}{chart_row}')
        chart_row += 16
        r += 2
    ws.column_dimensions['A'].width = 16
    for j in range(n + 3): ws.column_dimensions[L(c0 + j)].width = 12
    ws.freeze_panes = 'B4'

# 요약
sm = wb.create_sheet('요약', 1)
sm['A1'] = f'{NAME} · 시험 수준별 반 평균 변화'; sm['A1'].font = TITLE
sm['A2'] = '초록 글씨 = 각 시트의 반 평균 행에서 가져온 값'; sm['A2'].font = MUTED
for k, h in enumerate(['시험 수준', '영역', '첫 시험 평균', '마지막 시험 평균', '변화']): cell(sm, 4, 1 + k, h, BOLD, HDR)
for i, (g, name, wsn, ar, cf, cla, cd) in enumerate(summary):
    rr = 5 + i
    cell(sm, rr, 1, g, BOLD); cell(sm, rr, 2, name)
    cell(sm, rr, 3, f"='{wsn}'!{L(cf)}{ar}", GRN, fmt='0.0')
    cell(sm, rr, 4, f"='{wsn}'!{L(cla)}{ar}", GRN, fmt='0.0')
    cell(sm, rr, 5, f"=D{rr}-C{rr}", BLK, fmt='+0.0;-0.0;0.0')
for c, w in zip('ABCDE', [20, 10, 14, 16, 10]): sm.column_dimensions[c].width = w

INFO = {
 'A': ('2020–2022', '2020 1M-A 1차 모의고사 상위 4명', ['· 중2 1·2차는 성적표에 시험 수준 표기가 없습니다. 3차부터 "고2 모의고사"로 표기되어 있어 1·2차는 고1 모의고사로 분류했습니다.', '· 중1 1차는 듣기·독해 개수 기록이 없어 듣기·독해 그래프에서 빈칸입니다.', '· 나민하는 2023학년도 수능 날이 기말고사 기간이라 응시하지 않았습니다(빈칸).']),
 'B': ('2023–2025', '2023 1M-S부터 3년 동안 다닌 4명', ['· 중3 4차 고3 모의고사는 성적표에 아직 기록이 없습니다.', '· 중3 2차 구성관 총점은 2025년 성적표 두 장에 92점과 93점으로 다르게 적혀 있어, 듣기·독해 개수가 함께 적힌 92점을 사용했습니다.', '· 김의현·최리원(2023 1차만 응시), 백시우·신상훈(2025 합류)은 제외했습니다.']),
 'C': ('2025–2026', '2025 1M-A에서 2026 2M-A까지 다닌 6명', ['· 중1 1·2차는 총점만 기록된 학생이 있어 듣기·독해 칸이 비어 있습니다. 고2 모의고사는 2026년 중2 3차 한 번뿐이라 그래프에 점 하나만 찍힙니다.', '· 이서연 반은 현재 중2라 고3 모의고사는 아직 보지 않았습니다.', '· 한재승(2026 합류), 김은혜·구승준·김가민(2026 미등록)은 제외했습니다.'])}[CO]
lines = [(f'{NAME} 시험 수준별 성적 변화 ({INFO[0]})', TITLE),
 (f"대상: {INFO[1]} ({', '.join(STUDS)}) · Metta 영어 전문학원 성적표 사진에서 옮김", MUTED), ('', BLK), ('시트 구성', BOLD),
 ('· 요약: 시험 수준별·영역별 반 평균의 첫 시험 → 마지막 시험 변화', BLK)]
lines += [(f'· {ws}: {desc} · 듣기·독해·총점 표와 그래프', BLK) for ws, desc, _ in GROUPS]
lines += [('', BLK), ('읽는 법', BOLD), ('· 듣기 17문항, 독해 28문항 기준 맞은 개수 · 총점 100점', BLK),
 ('· 파란 글씨 = 성적표에서 옮긴 값(고치면 계산과 그래프가 바뀜) · 검은 글씨 = 계산식 · 초록 글씨 = 다른 시트 값', BLK), ('', BLK), ('참고', BOLD)]
lines += [(t, BLK) for t in INFO[2]]
for i, (t, f) in enumerate(lines, 1): info.cell(i, 1, t).font = f
info.column_dimensions['A'].width = 120
OUT = f"{NAME.replace(' ', '')}_시험수준별_성적.xlsx"; wb.save(OUT); print(OUT)
